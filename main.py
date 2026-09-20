import streamlit as st
import google.generativeai as genai
from pypdf import PdfReader
from pptx import Presentation

# إعداد الصفحة
st.set_page_config(page_title="MindSphere AI", page_icon="🤖", layout="wide")

# قراءة مفتاح الـ API تلقائياً من إعدادات المنصة السحابية (Secrets)
try:
    api_key = st.secrets["GOOGLE_API_KEY"]
    genai.configure(api_key=api_key)
except Exception:
    st.error("الرجاء ضبط مفتاح GOOGLE_API_KEY في إعدادات Secrets على منصة Streamlit.")
    st.stop()

st.title("🤖 MindSphere AI - تحليل الملفات الذكي")
st.write("أهلاً بكِ! قومي برفع ملف PDF أو PPTX وسيقوم الذكاء الاصطناعي بتحليله والإجابة على أسئلتكِ فوراً.")

# اختيار النموذج الذكي
model = genai.GenerativeModel("gemini-1.5-flash")

# رفع الملف
uploaded_file = st.file_uploader("اختر ملفاً (PDF أو PPTX)", type=["pdf", "pptx"])

file_text = ""

if uploaded_file is not None:
    file_extension = uploaded_file.name.split(".")[-1].lower()
    
    with st.spinner("جاري استخراج قراءة المحتوى..."):
        if file_extension == "pdf":
            reader = PdfReader(uploaded_file)
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    file_text += text + "\n"
        elif file_extension == "pptx":
            prs = Presentation(uploaded_file)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            file_text += paragraph.text + "\n"
                            
    st.success("تم قراءة الملف بنجاح! يمكنك الآن طرح أي سؤال حول محتواه.")

# صندوق لإدخال السؤال أو الاستفسار
user_question = st.text_input("ماذا تريدين أن تعرفي عن هذا الملف؟")

if user_question and file_text:
    with st.spinner("جاري تحليل المحتوى بواسطة Gemini..."):
        prompt = f"إليك محتوى الملف المستخرج:\n{question_context := file_text}\n\nسؤال المستخدم: {user_question}"
        response = model.generate_content(prompt)
        st.subheader("إجابة الذكاء الاصطناعي:")
        st.write(response.text)
elif user_question and not file_text:
    st.warning("الرجاء رفع ملف أولاً قبل طرح الأسئلة.")
